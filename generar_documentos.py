"""Genera los diagramas, los PDF de trabajo, el documento metodológico con formato del Manual y la página HTML.

Uso:
    python generar_documentos.py                 # diagramas SVG, PDF de trabajo y documento metodológico (PDF + DOCX)
    python generar_documentos.py pagina.html     # además, la página HTML del documento técnico en esa ruta

El documento metodológico tiene una sola versión: Documento-Metodologico.pdf y .docx, con portada (Manual,
anexo C-1), índice general (anexo C-2) y el texto desde el capítulo I. Los demás documentos salen en PDF de trabajo.

Requisitos: markdown, weasyprint y python-docx (pip), Pandoc para el .docx y, en Windows, Pango de MSYS2
(C:\\msys64\\mingw64\\bin), que el script detecta solo.
"""
import copy
import os
import re
import subprocess
import sys
from pathlib import Path

RAIZ = Path(__file__).parent
MSYS2_BIN = Path(r'C:\msys64\mingw64\bin')
if os.name == 'nt' and MSYS2_BIN.exists():
    os.environ.setdefault('WEASYPRINT_DLL_DIRECTORIES', str(MSYS2_BIN))

import markdown  # noqa: E402
import weasyprint  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_ROW_HEIGHT_RULE  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

sys.path.insert(0, str(RAIZ / 'diagramas'))
import generar as diagramas  # noqa: E402

EXTENSIONES = ['tables', 'fenced_code', 'sane_lists', 'toc']
DOCUMENTOS = ['Documento-Tecnico', 'Propuesta-Corta']  # PDF de trabajo
METODOLOGICO = 'Documento-Metodologico'
REFERENCIA_DOCX = RAIZ / 'plantillas' / 'referencia-manual.docx'
LOGO = RAIZ / 'plantillas' / 'logo-iupsm-azul.jpg'

# Datos de la portada (Manual, anexo C-1). El título va en pirámide invertida: cada línea, más corta que la anterior.
PORTADA = {
    'encabezado': ['REPÚBLICA BOLIVARIANA DE VENEZUELA',
                   'MINISTERIO DEL PODER POPULAR PARA LA', 'EDUCACIÓN UNIVERSITARIA,', 'CIENCIA Y TECNOLOGÍA',
                   'INSTITUTO UNIVERSITARIO POLITÉCNICO', '"SANTIAGO MARIÑO"', 'EXTENSIÓN CARACAS'],
    'titulo': ['Plataforma web institucional de alojamiento',
               'y despliegue continuo (PaaS) para el',
               'Politécnico Santiago Mariño'],
    'leyenda': ['Propuesta de Trabajo de Grado como requisito parcial para optar al Título de',
                'Ingeniero de Sistemas'],
    'datos': [('Autor (a)', 'José Javier Cuello Camargo'),
              ('Tutor (a)', 'Yancelis Alfredo Noguera Pinto'),
              ('Docente de la asignatura (a)', 'Carmen Díaz')],
    'lugar_fecha': 'Caracas, octubre 2026',
}
# Posiciones de la portada, en cm desde el borde superior del área de texto (la hoja tiene 3 cm de margen superior).
# Se tomaron de la guía del anexo C-1: el título queda a unos 12 cm del borde de la hoja, los datos a unos 20,5 cm
# y la ciudad con la fecha a unos 23,5 cm. El logo ocupa la franja izquierda del encabezado.
COLUMNA_LOGO, COLUMNA_DERECHA = 3.0, 2.0
FILAS_PORTADA = {'encabezado': 9.0, 'titulo': 8.5, 'datos': 3.0, 'lugar': 0.8}

# Índice general (Manual, anexo C-2). Las partes que aún no existen llevan "x" en lugar del número de página.
PRELIMINARES_INDICE = ['LISTA DE CUADROS', 'LISTA DE FIGURAS', 'RESUMEN', 'INTRODUCCIÓN']
ETIQUETAS_INDICE = {'Objetivos de la Investigación': 'Objetivos de la Investigación: general y específicos'}

CSS_PDF = """@page{size:Letter;margin:2cm}
body{font-family:'DejaVu Sans',sans-serif;font-size:10pt;line-height:1.45;color:#222}
h1{font-size:18pt;border-bottom:2px solid #333;padding-bottom:4px}
h2{font-size:14pt;margin-top:1.4em;border-bottom:1px solid #aaa}
h3{font-size:12pt;margin-top:1.2em}
h4{font-size:11pt;margin-top:1.1em;color:#333}
blockquote{border-left:3px solid #888;margin:0.6em 0;padding:0.2em 0.8em;background:#f4f4f4}
table{border-collapse:collapse;width:100%;font-size:8.5pt;margin:0.6em 0}
th,td{border:1px solid #999;padding:4px 6px;text-align:left;vertical-align:top}
th{background:#e8e8e8}
tr{page-break-inside:avoid}
code{font-family:'DejaVu Sans Mono',monospace;font-size:8.5pt}
hr{border:none;border-top:1px solid #ccc}
li{margin:0.15em 0}
p:has(> img){margin:0.8em 0 0.2em;page-break-inside:avoid}
img{width:100%}
p:has(> em:only-child){font-size:8.5pt;color:#555;margin-top:0}"""

# Formato del Manual del IUPSM, cap. V: carta; márgenes 4/3/3/3 cm; Arial 12; interlineado 1,5; sangría de
# 5 espacios; capítulos en página nueva a 5 cm; preliminares en romanos y cuerpo en arábigos, centrados abajo.
CSS_MANUAL = f"""
@page{{size:Letter;margin:3cm 3cm 3cm 4cm;
  @bottom-center{{content:counter(page);font:10pt Arial,'Liberation Sans',sans-serif;
    vertical-align:bottom;padding-bottom:1.5cm}}}}
@page portada{{@bottom-center{{content:none}}}}
@page preliminar{{@bottom-center{{content:counter(page,lower-roman)}}}}
body{{font-family:Arial,'Liberation Sans',sans-serif;font-size:12pt;line-height:1.5;color:#000}}
.portada{{page:portada;break-after:page;position:relative;height:21.9cm;text-align:center;line-height:1.15}}
.portada .logo{{position:absolute;left:0.3cm;top:0;width:2.5cm}}
.portada .encabezado{{position:absolute;top:0;left:{COLUMNA_LOGO}cm;right:{COLUMNA_DERECHA}cm;font-weight:bold;
  font-size:11pt}}
.portada .bloque-titulo{{position:absolute;top:{FILAS_PORTADA['encabezado']}cm;left:0;right:0}}
.portada .titulo{{font-weight:bold}}
.portada .leyenda{{font-size:11pt;margin-top:0.3cm}}
.portada .datos{{position:absolute;top:{FILAS_PORTADA['encabezado'] + FILAS_PORTADA['titulo']}cm;right:0;
  text-align:right;font-size:11pt}}
.portada .lugar{{position:absolute;
  top:{FILAS_PORTADA['encabezado'] + FILAS_PORTADA['titulo'] + FILAS_PORTADA['datos']}cm;left:0;right:0;
  font-size:11pt}}
.preliminar{{page:preliminar;break-after:page}}
.preliminar h2,h2.referencias{{text-align:center;font-size:12pt;font-weight:bold;margin:0 0 12pt}}
.indice ol{{list-style:none;padding:0;margin:0;line-height:1.3}}
.indice li{{display:flex;margin:0}}
.indice .pp{{justify-content:flex-end;font-weight:bold}}
.indice .pre{{font-weight:bold}}
.indice .rotulo{{font-weight:bold;margin-top:14pt}}
.indice .cap{{font-weight:bold;margin-top:12pt}}
.indice .ref{{margin-top:14pt}}
.indice .txt{{flex:1;min-width:0;overflow:hidden;white-space:nowrap}}
.indice .con-num .txt::after{{content:" ........................................................................................................................................"}}
.indice .num{{padding-left:4pt;min-width:0.8cm;text-align:right}}
p{{margin:0;text-indent:1.27cm;text-align:justify;orphans:2;widows:2}}
h2.capitulo{{break-before:page;padding-top:2cm;margin:0 0 24pt;text-align:center;font-size:12pt}}
h2.capitulo span{{display:block}}
h2.capitulo span+span{{margin-top:12pt}}
h2.referencias{{break-before:page}}
h3{{text-align:center;font-size:12pt;font-weight:bold;margin:24pt 0 12pt;break-after:avoid}}
h4{{text-align:center;font-size:12pt;font-weight:bold;font-style:italic;margin:24pt 0 12pt;break-after:avoid}}
h5{{text-align:left;font-size:12pt;font-weight:normal;font-style:italic;margin:18pt 0 6pt;break-after:avoid}}
ul,ol{{margin:0 0 0 1.27cm;padding-left:0.6cm}}
li{{text-align:justify}}
p.cuadro-num{{text-indent:0;margin-top:12pt;break-after:avoid}}
p.cuadro-titulo{{text-indent:0;break-after:avoid}}
table{{border-collapse:collapse;width:100%;font-size:11pt;line-height:1.15;margin:6pt 0 0;
  border-top:1px solid #000;border-bottom:1px solid #000;break-inside:avoid}}
th{{border-bottom:1px solid #000;text-align:left;padding:3pt 4pt}}
td{{text-align:left;vertical-align:top;padding:3pt 4pt}}
p.nota{{text-indent:0;font-size:10pt;line-height:1.15;margin:4pt 0 12pt}}
p.ref{{text-indent:-1.27cm;padding-left:1.27cm;line-height:1.15;margin-bottom:6pt;text-align:left}}
mark{{background:#fff0a0;font-size:10.5pt}}
"""

DARK = dict(fg='#dde6e8', muted='#9aabb1', line='#7f9299', surface='#1d282b', bg='#141c1f',
            accent='#4fb3c2', accentbg='#17363b', onaccent='#0e1618', ok='#5cc28a', bad='#f07a6e',
            oe1='#e0a948', oe2='#4fb3c2', oe3='#a58ce6')

CAPITULO = re.compile(r'^## (CAPÍTULO ([IVXL]+))\. (.+)$', re.M)


def generar_pdf(md, nombre):
    cuerpo = markdown.markdown(md, extensions=EXTENSIONES)
    html = (f'<!doctype html><html lang="es"><head><meta charset="utf-8"><style>{CSS_PDF}</style>'
            f'</head><body>{cuerpo}</body></html>')
    weasyprint.HTML(string=html, base_url=str(RAIZ)).write_pdf(RAIZ / f'{nombre}.pdf')


def secciones_desde_capitulo_i(md):
    return [s for s in re.split(r'(?m)^(?=## )', md[md.index('## CAPÍTULO I'):]) if s.strip()]


def es_pendiente(seccion):
    """Una sección está pendiente si, fuera de sus títulos, solo dice *Pendiente.*"""
    texto = re.sub(r'(?m)^#+ .*$', '', seccion).strip()
    return texto == '*Pendiente.*'


def extraer_entrega(md):
    """Toma desde el capítulo I hasta el final, sin capítulos pendientes, separadores ni notas internas."""
    md = ''.join(s for s in secciones_desde_capitulo_i(md) if not es_pendiente(s))
    md = re.sub(r'(?m)^---\s*$', '', md)
    md = re.sub(r'(?m)^\*Provisional[^\n]*\n', '', md)
    return md


def filas_indice(md, pagina):
    """Filas del índice general: (clase, texto, página). La página es None en los rótulos sin número."""
    filas = [('pre', p, 'x') for p in PRELIMINARES_INDICE]
    filas.append(('rotulo', 'CAPÍTULO', None))
    for seccion in secciones_desde_capitulo_i(md):
        titulo = seccion.splitlines()[0][3:].strip()
        pendiente = es_pendiente(seccion)
        if titulo == 'Referencias':
            filas.append(('pre ref', 'REFERENCIAS', pagina('REFERENCIAS') or 'x'))
            continue
        m = CAPITULO.match(seccion.splitlines()[0])
        filas.append(('cap', f'{m[2]}. {m[3]}', None))
        for sub in re.findall(r'(?m)^### (.+)$', seccion):
            filas.append(('sec', ETIQUETAS_INDICE.get(sub, sub), 'x' if pendiente else (pagina(sub) or 'x')))
    return filas


def portada_html():
    p = PORTADA
    datos = ''.join(f'<div><b>{k}:</b> {v}</div>' for k, v in p['datos'])
    return ('<section class="portada">'
            f'<img class="logo" src="{LOGO.relative_to(RAIZ).as_posix()}">'
            f'<div class="encabezado">{"<br>".join(p["encabezado"])}</div>'
            f'<div class="bloque-titulo"><div class="titulo">{"<br>".join(p["titulo"])}</div>'
            f'<div class="leyenda">{"<br>".join(p["leyenda"])}</div></div>'
            f'<div class="datos">{datos}</div>'
            f'<div class="lugar">{p["lugar_fecha"]}</div></section>')


def generar_metodologico_pdf(md):
    """Devuelve también las páginas de cada título, que el .docx usa como valor inicial de su índice."""
    cuerpo = markdown.markdown(extraer_entrega(md), extensions=EXTENSIONES)
    cuerpo = re.sub(r'<h2 id="([^"]+)">(CAPÍTULO [IVXL]+)\. (.*?)</h2>',
                    r'<h2 class="capitulo" id="\1"><span>\2</span><span>\3</span></h2>', cuerpo)
    i = cuerpo.find('<h2 id="referencias">')
    if i >= 0:
        cola = cuerpo[i:].replace('<h2 id="referencias">Referencias</h2>',
                                  '<h2 class="referencias" id="referencias">REFERENCIAS</h2>')
        cuerpo = cuerpo[:i] + cola.replace('<p>', '<p class="ref">')
    cuerpo = re.sub(r'<p><strong>(Cuadro \d+)</strong></p>\s*<p><strong>(.*?)</strong></p>',
                    r'<p class="cuadro-num"><strong>\1</strong></p><p class="cuadro-titulo"><strong>\2</strong></p>',
                    cuerpo)
    cuerpo = cuerpo.replace('<p><em>Nota.</em>', '<p class="nota"><em>Nota.</em>')
    cuerpo = re.sub(r'<code>(\[[^<]*\])</code>', r'<mark>\1</mark>', cuerpo)

    def documento(contenido):
        html = (f'<!doctype html><html lang="es"><head><meta charset="utf-8"><style>{CSS_MANUAL}</style>'
                f'</head><body>{contenido}</body></html>')
        return weasyprint.HTML(string=html, base_url=str(RAIZ)).render()

    # El cuerpo se pagina aparte para que empiece en 1; con sus anclas se arma el índice de las preliminares.
    doc_cuerpo = documento(cuerpo)
    pagina_de_ancla = {ancla: n for n, pag in enumerate(doc_cuerpo.pages, 1) for ancla in pag.anchors}
    paginas = {re.sub(r'<[^>]+>', '', texto): str(pagina_de_ancla[ident])
               for ident, texto in re.findall(r'<h[23][^>]*id="([^"]+)"[^>]*>(.*?)</h[23]>', cuerpo)
               if ident in pagina_de_ancla}

    filas = []
    for clase, texto, num in filas_indice(md, paginas.get):
        if num is None:
            filas.append(f'<li class="{clase}"><span class="txt">{texto}</span></li>')
        else:
            filas.append(f'<li class="{clase} con-num"><span class="txt">{texto}</span>'
                         f'<span class="num">{num}</span></li>')
    indice = ('<section class="preliminar indice"><h2>ÍNDICE GENERAL</h2>'
              f'<ol><li class="pp"><span>Pp.</span></li>{"".join(filas)}</ol></section>')
    doc_prelim = documento(portada_html() + indice)

    destino = RAIZ / f'{METODOLOGICO}.pdf'
    doc_prelim.copy(doc_prelim.pages + doc_cuerpo.pages).write_pdf(destino)
    return destino, paginas


# --- Documento metodológico en Word -------------------------------------------------------------------------------

def _parrafo(contenedor, texto='', estilo=None, negrita=False):
    par = contenedor.add_paragraph(style=estilo)
    if texto:
        par.add_run(texto).bold = negrita or None
    return par


def _sin_bordes(tabla):
    bordes = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:val'), 'nil')
        bordes.append(el)
    tabla._tbl.tblPr.append(bordes)


def _campo_pageref(par, marcador, valor):
    """Número de página como campo de Word: se actualiza al abrir el documento; mientras, muestra el del PDF."""
    for tipo, texto in (('begin', None), ('instr', f' PAGEREF {marcador} \\h '), ('separate', None),
                        ('texto', valor), ('end', None)):
        run = par.add_run()
        if tipo == 'instr':
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = texto
            run._r.append(el)
        elif tipo == 'texto':
            run.text = texto
        else:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tipo)
            run._r.append(el)


def _portada_docx(doc):
    p = PORTADA
    ancho = 21.59 - 4 - 3
    tabla = doc.add_table(rows=4, cols=3)
    tabla.autofit = False
    _sin_bordes(tabla)
    anchos = [COLUMNA_LOGO, ancho - COLUMNA_LOGO - COLUMNA_DERECHA, COLUMNA_DERECHA]
    for fila in tabla.rows:
        for celda, w in zip(fila.cells, anchos):
            celda.width = Cm(w)
    for fila, alto in zip(tabla.rows, FILAS_PORTADA.values()):
        fila.height, fila.height_rule = Cm(alto), WD_ROW_HEIGHT_RULE.EXACTLY

    logo = tabla.cell(0, 0).paragraphs[0]
    logo.paragraph_format.left_indent = Cm(0.1)
    logo.add_run().add_picture(str(LOGO), width=Cm(2.5))
    celda = tabla.cell(0, 1)
    celda.paragraphs[0].style = 'Portada Encabezado'
    celda.paragraphs[0].add_run(p['encabezado'][0])
    for linea in p['encabezado'][1:]:
        _parrafo(celda, linea, 'Portada Encabezado')

    celda = tabla.cell(1, 0).merge(tabla.cell(1, 2))
    celda.paragraphs[0].style = 'Portada Titulo'
    celda.paragraphs[0].add_run(p['titulo'][0])
    for linea in p['titulo'][1:]:
        _parrafo(celda, linea, 'Portada Titulo')
    for n, linea in enumerate(p['leyenda']):
        par = _parrafo(celda, linea, 'Portada Leyenda')
        if n == 0:
            par.paragraph_format.space_before = Cm(0.3)

    celda = tabla.cell(2, 0).merge(tabla.cell(2, 2))
    for n, (k, v) in enumerate(p['datos']):
        par = celda.paragraphs[0] if n == 0 else celda.add_paragraph()
        par.style = 'Portada Datos'
        par.add_run(f'{k}:').bold = True
        par.add_run(f' {v}')

    celda = tabla.cell(3, 0).merge(tabla.cell(3, 2))
    celda.paragraphs[0].style = 'Portada Lugar'
    celda.paragraphs[0].add_run(p['lugar_fecha'])
    return tabla


def _indice_docx(doc, md, paginas, marcadores):
    ancho = Cm(21.59 - 4 - 3)
    elementos = []

    def fila(texto, clase, num):
        par = doc.add_paragraph(style='Compact')
        pf = par.paragraph_format
        pf.first_line_indent, pf.line_spacing, pf.space_after = Cm(0), 1.15, Pt(0)
        if clase in ('rotulo', 'cap') or 'ref' in clase:
            pf.space_before = Pt(12)
        run = par.add_run(texto)
        run.bold = clase != 'sec' or None
        if num is not None:
            pf.tab_stops.add_tab_stop(ancho, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            par.add_run('\t')
            clave = 'REFERENCIAS' if texto == 'REFERENCIAS' else next(
                (k for k, v in ETIQUETAS_INDICE.items() if v == texto), texto)
            if num != 'x' and clave in marcadores:
                _campo_pageref(par, marcadores[clave], num)
            else:
                par.add_run(num)
        elementos.append(par)

    titulo = doc.add_paragraph(style='Compact')
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.page_break_before = True
    titulo.paragraph_format.space_after = Pt(12)
    titulo.add_run('ÍNDICE GENERAL').bold = True
    elementos.append(titulo)
    pp = doc.add_paragraph(style='Compact')
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pp.add_run('Pp.').bold = True
    elementos.append(pp)
    for clase, texto, num in filas_indice(md, paginas.get):
        fila(texto, clase, num)
    return elementos


def _numeracion(sect_pr, formato, inicio=1, portada=False):
    for viejo in sect_pr.findall(qn('w:pgNumType')) + sect_pr.findall(qn('w:titlePg')):
        sect_pr.remove(viejo)
    num = OxmlElement('w:pgNumType')
    num.set(qn('w:fmt'), formato)
    num.set(qn('w:start'), str(inicio))
    # El orden de los hijos de sectPr importa: pgNumType va después de pgMar/pgSz y antes de cols.
    ancla = sect_pr.find(qn('w:cols'))
    if ancla is not None:
        ancla.addprevious(num)
    else:
        sect_pr.append(num)
    if portada:
        titulo = OxmlElement('w:titlePg')
        posterior = sect_pr.find(qn('w:docGrid'))
        if posterior is not None:
            posterior.addprevious(titulo)
        else:
            sect_pr.append(titulo)


def generar_metodologico_docx(md, paginas):
    def bloque(estilo, lineas):
        return f'::: {{custom-style="{estilo}"}}\n' + '\\\n'.join(lineas) + '\n:::\n\n'

    cuerpo = extraer_entrega(md)
    cuerpo = CAPITULO.sub(lambda m: f'## {m[1]}\n\n::: {{custom-style="Nombre de Capítulo"}}\n{m[3]}\n:::', cuerpo)
    cuerpo = re.sub(r'(?m)^\*\*(Cuadro \d+)\*\*\n\n\*\*(.+?)\*\*$',
                    lambda m: bloque('Cuadro Numero', [f'**{m[1]}**']) + bloque('Cuadro Titulo', [f'**{m[2]}**']),
                    cuerpo)
    cuerpo = re.sub(r'(?m)^\*Nota\.\*(.*)$', lambda m: bloque('Nota de Cuadro', [f'*Nota.*{m[1]}']), cuerpo)
    i = cuerpo.find('## Referencias')
    if i >= 0:
        refs = [r.strip() for r in cuerpo[i + len('## Referencias'):].split('\n\n') if r.strip()]
        cuerpo = cuerpo[:i] + '## REFERENCIAS\n\n' + ''.join(bloque('Referencia', [r]) for r in refs)

    destino = RAIZ / f'{METODOLOGICO}.docx'
    subprocess.run(['pandoc', '-f', 'markdown', '-t', 'docx', '--reference-doc', str(REFERENCIA_DOCX),
                    '-o', str(destino)], input=cuerpo, text=True, encoding='utf-8', check=True)

    # Pandoc arma el cuerpo; la portada y el índice se agregan delante, en una sección con numeración romana.
    doc = Document(destino)
    cuerpo_xml = doc.element.body
    primero = cuerpo_xml[0]
    # Marcadores ocultos (nombre con guion bajo inicial) dentro de cada título, para los números del índice.
    marcadores = {}
    for n, par in enumerate(p for p in doc.paragraphs if p.style.name.startswith('Heading')):
        nombre = f'_Indice{n}'
        inicio, fin = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
        inicio.set(qn('w:id'), str(9000 + n))
        inicio.set(qn('w:name'), nombre)
        fin.set(qn('w:id'), str(9000 + n))
        par._p.get_or_add_pPr().addnext(inicio)
        par._p.append(fin)
        marcadores[par.text.strip()] = nombre

    tabla = _portada_docx(doc)
    elementos = [tabla._tbl] + [e._p for e in _indice_docx(doc, md, paginas, marcadores)]
    fin_prelim = doc.add_paragraph(style='Compact')
    elementos.append(fin_prelim._p)
    for el in elementos:
        primero.addprevious(el)

    sect_final = cuerpo_xml.find(qn('w:sectPr'))
    sect_prelim = copy.deepcopy(sect_final)
    _numeracion(sect_prelim, 'lowerRoman', portada=True)
    fin_prelim._p.get_or_add_pPr().append(sect_prelim)
    _numeracion(sect_final, 'decimal')

    actualizar = OxmlElement('w:updateFields')
    actualizar.set(qn('w:val'), 'true')
    doc.settings.element.append(actualizar)
    doc.save(destino)
    return destino


def generar_pagina(md, destino):
    svgs = diagramas.generar('variables')
    cuerpo = markdown.markdown(md, extensions=EXTENSIONES)
    cuerpo = re.sub(r'<h1[^>]*>.*?</h1>\s*', '', cuerpo, count=1)

    def figura(m):
        return (f'<figure><div class="lienzo">{svgs[m.group(1)]}</div>'
                f'<figcaption>{m.group(2)}</figcaption></figure>')
    cuerpo = re.sub(r'<p><img alt="[^"]*" src="diagramas/([\w-]+)\.svg" ?/?></p>\s*<p><em>(.*?)</em></p>',
                    figura, cuerpo, flags=re.S)
    cuerpo = re.sub(r'(<table>.*?</table>)', r'<div class="tabla">\1</div>', cuerpo, flags=re.S)
    cuerpo = cuerpo.replace('<li>[ ] ', '<li class="tarea">').replace('<li>\n<p>[ ] ', '<li class="tarea">\n<p>')

    indice = ''.join(
        f'<li class="n{nivel}"><a href="#{ident}">{re.sub("<[^>]+>", "", texto)}</a></li>'
        for nivel, ident, texto in re.findall(r'<h([23]) id="([^"]+)">(.*?)</h\1>', cuerpo))

    claro = ''.join(f'--d-{k}:{v};' for k, v in diagramas.ESTATICO.items() if k != 'font')
    oscuro = ''.join(f'--d-{k}:{v};' for k, v in DARK.items())
    plantilla = (RAIZ / 'diagramas' / 'plantilla.html').read_text(encoding='utf-8')
    pagina = (plantilla.replace('/*D-CLARO*/', claro).replace('/*D-OSCURO*/', oscuro)
              .replace('<!--INDICE-->', indice).replace('<!--CUERPO-->', cuerpo))
    Path(destino).write_text(pagina, encoding='utf-8')


if __name__ == '__main__':
    for nombre, svg in diagramas.generar().items():
        (RAIZ / 'diagramas' / f'{nombre}.svg').write_text(svg, encoding='utf-8')
    for nombre in DOCUMENTOS:
        generar_pdf((RAIZ / f'{nombre}.md').read_text(encoding='utf-8'), nombre)
    metodologico = (RAIZ / f'{METODOLOGICO}.md').read_text(encoding='utf-8')
    pdf, paginas = generar_metodologico_pdf(metodologico)
    print('Documento metodológico:', pdf)
    print('Documento metodológico:', generar_metodologico_docx(metodologico, paginas))
    if len(sys.argv) > 1:
        generar_pagina((RAIZ / 'Documento-Tecnico.md').read_text(encoding='utf-8'), sys.argv[1])
