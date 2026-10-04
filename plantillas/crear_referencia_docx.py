"""Crea plantillas/referencia-manual.docx: la plantilla de estilos que Pandoc usa para el .docx entregable.

Aplica el formato del Manual del IUPSM (5ª ed., cap. V): papel carta; márgenes de 4 cm a la izquierda y 3 cm en
los demás lados; Arial 12; interlineado 1,5; sangría de 5 espacios; títulos de capítulo centrados en página nueva;
número de página centrado abajo.

Uso:
    python plantillas/crear_referencia_docx.py
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DESTINO = Path(__file__).parent / 'referencia-manual.docx'
FUENTE = 'Arial'
SANGRIA = Cm(1.27)  # cinco espacios


def fuente(style, size=12, bold=None, italic=None, caps=None):
    f = style.font
    f.name, f.size, f.color.rgb = FUENTE, Pt(size), RGBColor(0, 0, 0)
    f.bold, f.italic, f.all_caps = bold, italic, caps
    rfonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        rfonts.attrib.pop(qn(attr), None)
    rfonts.set(qn('w:cs'), FUENTE)


def parrafo(style, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=SANGRIA, antes=0, despues=0, interlineado=1.5,
            salto=False, con_siguiente=False):
    pf = style.paragraph_format
    pf.alignment, pf.first_line_indent = align, indent
    pf.space_before, pf.space_after = Pt(antes), Pt(despues)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = interlineado
    pf.page_break_before, pf.keep_with_next = salto, con_siguiente


def estilo(doc, nombre, tipo=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[nombre]
    except KeyError:
        s = doc.styles.add_style(nombre, tipo)
        if tipo == WD_STYLE_TYPE.PARAGRAPH:
            s.base_style = doc.styles['Normal']
        return s


def numero_de_pagina(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.name, run.font.size = FUENTE, Pt(10)
    for tag, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
        run._r.append(el)


def main():
    base = DESTINO.with_name('_pandoc-default.docx')
    subprocess.run(['pandoc', '-o', str(base), '--print-default-data-file', 'reference.docx'], check=True)
    doc = Document(base)
    base.unlink()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = Cm(4)
    sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(3)
    sec.footer_distance = Cm(1.5)
    numero_de_pagina(sec)

    for nombre in ('Normal', 'Body Text', 'First Paragraph'):
        fuente(doc.styles[nombre])
        parrafo(doc.styles[nombre])
    fuente(doc.styles['Compact'])
    parrafo(doc.styles['Compact'], align=WD_ALIGN_PARAGRAPH.LEFT, indent=None)

    # Capítulo: "CAPÍTULO I" en página nueva, a 5 cm del borde (3 cm de margen + 2 cm).
    h2 = doc.styles['Heading 2']
    fuente(h2, bold=True, caps=True)
    parrafo(h2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=None, antes=57, despues=12, salto=True, con_siguiente=True)
    nombre_cap = estilo(doc, 'Nombre de Capítulo')
    fuente(nombre_cap, bold=True, caps=True)
    parrafo(nombre_cap, align=WD_ALIGN_PARAGRAPH.CENTER, indent=None, despues=24, con_siguiente=True)
    # Subtítulos sin numerar (Manual, figura 5).
    niveles = {'Heading 3': dict(bold=True), 'Heading 4': dict(bold=True, italic=True),
               'Heading 5': dict(bold=False, italic=True)}
    for nombre, f in niveles.items():
        fuente(doc.styles[nombre], **f)
        align = WD_ALIGN_PARAGRAPH.LEFT if nombre == 'Heading 5' else WD_ALIGN_PARAGRAPH.CENTER
        parrafo(doc.styles[nombre], align=align, indent=None, antes=24, despues=12, con_siguiente=True)

    for nombre, kw in {
        'Portada Encabezado': dict(bold=True, size=11),
        'Portada Titulo': dict(bold=True),
        'Portada Leyenda': dict(size=11),
        'Portada Datos': dict(size=11),
        'Portada Lugar': dict(size=11),
    }.items():
        s = estilo(doc, nombre)
        fuente(s, **kw)
        align = WD_ALIGN_PARAGRAPH.RIGHT if nombre == 'Portada Datos' else WD_ALIGN_PARAGRAPH.CENTER
        parrafo(s, align=align, indent=None, interlineado=1.0)

    ref = estilo(doc, 'Referencia')
    fuente(ref)
    parrafo(ref, align=WD_ALIGN_PARAGRAPH.LEFT, indent=-SANGRIA, despues=6, interlineado=1.0)
    ref.paragraph_format.left_indent = SANGRIA

    # Los marcadores [PENDIENTE…] van en `código`: se resaltan para que no pasen desapercibidos.
    verb = doc.styles['Verbatim Char']
    fuente(verb, size=11)
    verb.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.save(DESTINO)
    print(f'Plantilla creada: {DESTINO}')


if __name__ == '__main__':
    main()
